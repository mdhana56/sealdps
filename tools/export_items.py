"""Export semua item Seal Online (etc.SPAK -> item.edt + ItemString.edt) ke .xlsx.

Baca file game langsung (offline, cuma file di disk). Format:
- .SPAK = zip ZipCrypto, password dari versi di komentar zip (lihat src/DpsMeterUI/SkillTable.cs).
- .edt  = XOR stream LCG (seed 0x11CFD).
- item.edt       = "Seal Online Data v13": header 64 + count + field_count, baris 92 x int32.
- ItemString.edt = "Seal Online Data v13": baris = int32 id + nama[260] + deskripsi[512].
Nama kolom item.edt sebagian dari project open-source "unsealed" (hasil korelasi, belum
resmi) - lihat sheet "Kolom".

Usage: python tools/export_items.py [output.xlsx]
"""
import re
import struct
import sys
import zipfile
import zlib
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SPAK = Path(r"C:\Program Files (x86)\SealOnline\etc\etc.SPAK")
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "exports" / "SealOnline_Items.xlsx"
TEMPLATES = ("#!#0%x&&!!", "^^&!@%X&&*", "#$#$&*%x!!@", "!@####%x*@#@")

# ---- SPAK / EDT --------------------------------------------------------------

CRC = []
for n in range(256):
    c = n
    for _ in range(8):
        c = (c >> 1) ^ 0xEDB88320 if c & 1 else c >> 1
    CRC.append(c)


def spak_password(version: int) -> bytes:
    buf = (b"Pass" + str(99999999 - version).encode()).ljust(260, b"\x00")
    return (TEMPLATES[version % 4] % (zlib.crc32(buf) & 0xFFFFFFFF)).encode()


def zipcrypto_decrypt(pw: bytes, data: bytes) -> bytes:
    k0, k1, k2 = 0x12345678, 0x23456789, 0x34567890

    def upd(ch):
        nonlocal k0, k1, k2
        k0 = (k0 >> 8) ^ CRC[(k0 ^ ch) & 0xFF]
        k1 = ((k1 + (k0 & 0xFF)) * 134775813 + 1) & 0xFFFFFFFF
        k2 = (k2 >> 8) ^ CRC[(k2 ^ (k1 >> 24)) & 0xFF]

    for b in pw:
        upd(b)
    out = bytearray(len(data))
    for i, c in enumerate(data):
        k = k2 | 2
        c ^= ((k * (k ^ 1)) >> 8) & 0xFF
        upd(c)
        out[i] = c
    return bytes(out)


def edt_decode(data: bytes) -> bytes:
    out = bytearray(len(data))
    seed = 0x11CFD
    for i, c in enumerate(data):
        out[i] = c ^ ((seed >> 8) & 0xFF)
        seed = ((seed + (c if c < 0x80 else c - 0x100)) * 52845 + 22719) & 0xFFFF
    return bytes(out)


def read_entry(zf: zipfile.ZipFile, fp, pw: bytes, name: str) -> bytes:
    info = zf.getinfo(name)
    fp.seek(info.header_offset)
    h = struct.unpack("<IHHHHHIIIHH", fp.read(30))
    fp.seek(info.header_offset + 30 + h[9] + h[10])
    raw = fp.read(info.compress_size)
    if info.flag_bits & 1:
        raw = zipcrypto_decrypt(pw, raw)[12:]
    out = zlib.decompress(raw, -15) if info.compress_type == 8 else raw
    if zlib.crc32(out) & 0xFFFFFFFF != info.CRC:
        raise ValueError(f"CRC {name} tidak cocok")
    return out


# ---- Kolom item.edt ------------------------------------------------------------
# (nama, tingkat keyakinan, catatan). Kolom yang tidak ada di sini tampil sebagai "cN".
KNOWN = {
    1: ("Tipe", "dugaan", "Hellscream: 17=Helmet, 16=Top, 26=Bottom, 20=Shoes"),
    3: ("Level dasar?", "dugaan", "Einhorn helmet XG = 252 tapi tooltip Level Requirement 261 - belum pasti"),
    5: ("Fame Requirement", "terverifikasi", "Einhorn helmet XG = 112.256 = tooltip"),
    9: ("Damage Increase %", "terverifikasi", "Nilai +0. Einhorn Normal helmet XG = 4 (in-game +0: 4%), top XG = 3 (+12: 6%). Refine +12 menambah +3"),
    20: ("Damage Decrease %", "terverifikasi", "Nilai +0. Einhorn Normal helmet XG = 13 (in-game +0: 13%), top XG = 13 (+12: 16%). Refine +12 menambah +3"),
    6: ("Damage", "terverifikasi", "Nilai +0. Einhorn helmet XG = 10, tooltip +12 = 70"),
    12: ("Magic Power", "terverifikasi", "Nilai +0. Einhorn helmet XG = 10, tooltip +12 = 70"),
    15: ("Min STA", "unsealed", ""),
    17: ("Defense", "terverifikasi", "Nilai +0. Einhorn helmet XG = 395, tooltip +12 = 515"),
    18: ("STA+/refine", "unsealed (float?)", "unsealed bilang float, tapi Hellscream berisi int 1 - arti belum pasti"),
    23: ("Attack Speed", "terverifikasi", "Einhorn helmet XG = 18 = tooltip"),
    25: ("Accuracy", "terverifikasi", "Einhorn helmet XG = 5 = tooltip"),
    27: ("Critical Rate", "terverifikasi", "Einhorn helmet XG = 5 = tooltip"),
    29: ("Evasion", "unsealed", ""),
    31: ("Moving Speed", "terverifikasi", "Einhorn helmet XG = 4 = tooltip (unsealed menyebut Luck - salah)"),

    33: ("Set ID", "unsealed", "Hellscream: 1126 biasa, 1127 .G, 1128 .DG, 1129 .XG"),
    35: ("Buy Price", "unsealed", ""),
    37: ("HP increase", "terverifikasi", "Einhorn helmet XG = 27 = tooltip (di potion: HP restore)"),
    39: ("AP increase", "terverifikasi", "Einhorn helmet XG = 27 = tooltip (di potion: AP restore)"),
    42: ("Cooldown", "unsealed", ""),
    51: ("Upgrade ke ID", "dugaan", "Hellscream Top -> Top.G -> Top.DG -> Top.XG"),
    55: ("Refine AC item", "unsealed", ""),
    56: ("Refine A item", "unsealed", ""),
    57: ("Refine C item", "unsealed", ""),
    60: ("Attack Range", "unsealed", ""),
    61: ("Min STR", "unsealed", ""),
    63: ("Min AGI", "unsealed", ""),
    65: ("Min INT", "unsealed", ""),
    69: ("Min WIS", "unsealed", ""),
    75: ("Model ID", "unsealed", ""),
    76: ("Icon ID", "unsealed", ""),
}


def decode_text(raw: bytes) -> str:
    """Teks item campuran: nama English = ASCII, sebagian teks Korea tersimpan UTF-8,
    sebagian CP949. Coba UTF-8 dulu, lalu CP949; byte yang tetap tidak valid dibuang
    (bukan diganti U+FFFD)."""
    for enc in ("utf-8", "cp949"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            pass
    return raw.decode("cp949", "ignore")


def clean_text(s: str) -> str:
    # Kode format tooltip game selalu "#huruf#": #N# = baris baru, #B# #X# #R# #G# = warna.
    # (Dulu pakai "#[A-Z]" -> huruf pertama kata sesudahnya ikut kehapus: "##dditional".)
    s = s.replace("#N#", "\n")
    return re.sub(r"#[A-Z]#", "", s).strip()


def main():
    zf = zipfile.ZipFile(SPAK)
    version = int(re.search(rb"Seal Online Zip v(\d+)", zf.comment).group(1))
    pw = spak_password(version)
    with open(SPAK, "rb") as fp:
        print("decode item.edt ...")
        items = edt_decode(read_entry(zf, fp, pw, "item.edt"))
        print("decode ItemString.edt ...")
        strs = edt_decode(read_entry(zf, fp, pw, "ItemString.edt"))

        print("decode set_opt.edt ...")
        set_opt = edt_decode(read_entry(zf, fp, pw, "set_opt.edt"))

    n_items, n_fields = struct.unpack_from("<ii", items, 64)
    width = (len(items) - 72) // n_items
    assert width == n_fields * 4, (width, n_fields)
    n_str = struct.unpack_from("<i", strs, 64)[0]
    str_width = (len(strs) - 72) // n_str

    names = {}
    for r in range(n_str):
        b = 72 + r * str_width
        iid = struct.unpack_from("<i", strs, b)[0]
        nm = decode_text(strs[b + 4:b + 264].split(b"\0")[0]).strip()
        desc = decode_text(strs[b + 264:b + str_width].split(b"\0")[0])
        names[iid] = (nm, clean_text(desc))

    rows = [struct.unpack_from(f"<{n_fields}i", items, 72 + r * width) for r in range(n_items)]
    # Kolom mentah yang semua isinya 0 dibuang biar sheet nggak kebanyakan kolom.
    used = [c for c in range(1, n_fields) if c in KNOWN or any(v[c] for v in rows)]
    named = [c for c in used if c in KNOWN]
    raw = [c for c in used if c not in KNOWN]

    wb = Workbook(write_only=True)
    font = Font(name="Arial", size=10)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="305496")
    raw_fill = PatternFill("solid", fgColor="7F7F7F")

    from openpyxl.cell import WriteOnlyCell

    ws = wb.create_sheet("Items")
    ws.freeze_panes = "C2"
    headers = [("ID", head_fill), ("Nama", head_fill)]
    headers += [(f"{KNOWN[c][0]} [c{c}]", head_fill) for c in named]
    headers += [("Deskripsi", head_fill)]
    headers += [(f"c{c}", raw_fill) for c in raw]
    widths = [8, 42] + [14] * len(named) + [60] + [8] * len(raw)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n_items + 1}"

    head_row = []
    for text, fill in headers:
        cell = WriteOnlyCell(ws, value=text)
        cell.font, cell.fill = head_font, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        head_row.append(cell)
    ws.append(head_row)

    def cell(v, wrap=False):
        c = WriteOnlyCell(ws, value=v)
        c.font = font
        if isinstance(v, int) and abs(v) >= 1000:
            c.number_format = "#,##0"
        if wrap:
            c.alignment = Alignment(wrap_text=False, vertical="top")
        return c

    for v in rows:
        nm, desc = names.get(v[0], ("", ""))
        ws.append([cell(v[0]), cell(nm)] + [cell(v[c]) for c in named] + [cell(desc, wrap=True)] + [cell(v[c]) for c in raw])

    # Sheet set bonus: set_opt.edt = [set ID][jumlah bagian dipakai][15 nilai bonus].
    # Arti b1..b15 BELUM diketahui - dicocokkan nanti dengan tooltip in-game.
    so_count, so_fields = struct.unpack_from("<ii", set_opt, 64)
    so_width = (len(set_opt) - 72) // so_count
    set_rows = [struct.unpack_from(f"<{so_fields}i", set_opt, 72 + r * so_width) for r in range(so_count)]
    members = {}
    set_col = 33  # kolom Set ID di item.edt
    for v in rows:
        if v[set_col]:
            members.setdefault(v[set_col], []).append(names.get(v[0], ("", ""))[0])
    ss = wb.create_sheet("Set Bonus")
    ss.freeze_panes = "C2"
    # Terverifikasi dari tooltip in-game set Einhorn (527) 4 bagian: Damage 155, MPW 155, Defense 270.
    # b4..b15 terverifikasi dari tooltip set Hellscream (Magic) XG (1133) & Dragon Inferno
    # Swordmaster (839/1159). b5/b6/b7 belum diketahui, b11/b12 selalu 0.
    set_labels = {1: "Damage [b1]", 2: "Magic Power [b2]", 3: "Defense [b3]", 4: "Attack Speed [b4]",
                  8: "Moving Speed [b8]", 9: "HP % [b9]", 10: "AP % [b10]",
                  13: "Damage Increase % [b13]", 14: "Damage Decrease % [b14]", 15: "Dungeon Dmg Increase [b15]"}
    s_head = ["Set ID", "Jumlah bagian"] + [set_labels.get(k, f"b{k}") for k in range(1, so_fields - 1)] + ["Item di set ini"]
    for i, w in enumerate([8, 12] + [9] * (so_fields - 2) + [90], 1):
        ss.column_dimensions[get_column_letter(i)].width = w
    ss.auto_filter.ref = f"A1:{get_column_letter(len(s_head))}{so_count + 1}"
    hr = []
    for t in s_head:
        c = WriteOnlyCell(ss, value=t)
        c.font, c.fill = head_font, head_fill
        hr.append(c)
    ss.append(hr)

    def scell(v):
        c = WriteOnlyCell(ss, value=v)
        c.font = font
        if isinstance(v, int) and abs(v) >= 1000:
            c.number_format = "#,##0"
        return c

    for r in set_rows:
        if r[0] == 0:
            continue
        mem = members.get(r[0], [])
        ss.append([scell(x) for x in r] + [scell(", ".join(mem[:8]) + (f" (+{len(mem) - 8} lagi)" if len(mem) > 8 else ""))])

    # Sheet penjelasan kolom
    ks = wb.create_sheet("Kolom")
    for col, w in zip("ABCDE", (10, 22, 22, 70, 30)):
        ks.column_dimensions[col].width = w
    kh = []
    for t in ("Kolom", "Nama", "Keyakinan", "Catatan", "Contoh (Hellscream Vanguard Top, ID 35827)"):
        c = WriteOnlyCell(ks, value=t)
        c.font, c.fill = head_font, head_fill
        kh.append(c)
    ks.append(kh)
    example = next((v for v in rows if v[0] == 35827), None)

    def kcell(v):
        c = WriteOnlyCell(ks, value=v)
        c.font = font
        return c

    for c in used:
        name, conf, note = KNOWN.get(c, (f"c{c}", "belum diketahui", ""))
        ks.append([kcell(f"c{c}"), kcell(name), kcell(conf), kcell(note), kcell(example[c] if example else None)])
    ks.append([])
    for line in (
        f"Sumber: {SPAK} -> item.edt ({n_items} baris x {n_fields} kolom int32) + ItemString.edt (nama/deskripsi).",
        "Keyakinan 'unsealed' = nama kolom dari project open-source unsealed (hasil korelasi, bukan dokumentasi resmi).",
        "'dugaan' = disimpulkan dari data Hellscream; 'belum diketahui' = arti kolom belum dicari.",
        "Kolom mentah yang isinya 0 untuk semua item tidak ditampilkan.",
        "Sheet 'Set Bonus' = set_opt.edt: [Set ID][jumlah bagian][b1..b15]. Terverifikasi dari tooltip in-game"
        " (set Einhorn 527, Hellscream Magic XG 1133, Dragon Inferno 839): b1 Damage, b2 MPW, b3 Defense,"
        " b4 Attack Speed, b8 Moving Speed, b9 HP %, b10 AP %, b13 Damage Increase %, b14 Damage Decrease %,"
        " b15 Dungeon Dmg Increase. b5/b6/b7 belum diketahui.",
        "Semua nilai = item +0 (belum di-refine). Refine menambah sebagian stat (mis. +12 = +3 Damage Increase/Decrease).",
        "Cocokkan dengan tooltip in-game untuk memastikan arti kolom.",
    ):
        ks.append([kcell(line)])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{n_items} item, {len(named)} kolom bernama + {len(raw)} kolom mentah -> {OUT}")


if __name__ == "__main__":
    main()
